"""デモ用アンケート回答データを生成する。

共有テストデータ (Mirai_Questionnaire_TestData.xlsx) と同じ列構造を使う。
実在の生徒名・実回答は含めない。
"""

from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# 共有テストデータと同じ列名（typo の study_daily_mitutes もそのまま使う）
HEADERS = [
    "timestamp",
    "student_name",
    "student_id",
    "grade",
    "juku_satisfaction",
    "juku_security",
    "juku_fun",
    "juku_clarity",
    "juku_question",
    "juku_motivation",
    "juku_homework",
    "juku_growth",
    "juku_continuation",
    "home_plan_achievement",
    "home_study_environment",
    "home_pace_drop",
    "home_problem_solving",
    "home_motivation",
    "home_growth",
    "home_continuation",
    "school_understanding",
    "school_security",
    "school_fun",
    "school_clarity",
    "school_question",
    "school_class_motivation",
    "school_class_satisfaction",
    "school_test_motivation",
    "school_test_satisfaction",
    "school_growth",
    "school_test_expectation",
    "school_continuation",
    "study_fun",
    "study_goal",
    "study_improvement_motivation",
    "study_daily_mitutes",
    "study_concerns",
    "study_next_month_contents",
]

SCORE_FIELDS = [
    "juku_satisfaction",
    "juku_security",
    "juku_fun",
    "juku_clarity",
    "juku_question",
    "juku_motivation",
    "juku_homework",
    "juku_growth",
    "juku_continuation",
    "home_plan_achievement",
    "home_study_environment",
    "home_pace_drop",
    "home_problem_solving",
    "home_motivation",
    "home_growth",
    "home_continuation",
    "school_understanding",
    "school_security",
    "school_fun",
    "school_clarity",
    "school_question",
    "school_class_motivation",
    "school_class_satisfaction",
    "school_test_motivation",
    "school_test_satisfaction",
    "school_growth",
    "school_test_expectation",
    "school_continuation",
    "study_fun",
    "study_goal",
    "study_improvement_motivation",
]

# 共有テストデータで確認できた選択肢のみを使う
STUDY_MINUTES = {
    "10_30": "10～30分くらい",
    "30_60": "30～60分くらい",
}

CONCERNS = {
    "calc": "計算ミスや文字の書き間違いが多い",
    "applied": "応用問題が難しい",
    "score": "テストの点数が伸びない",
    "motivation": "やる気が出ないことが多い",
    "what": "何を勉強すればよいのか分からない",
    "how": "勉強のやり方が分からない",
}

NEXT_MONTH = {
    "reading": "国語の文章問題",
    "kanji": "漢字",
    "math": "算数・数学",
    "english": "英語",
    "science": "理科",
    "social": "社会",
    "hard": "難しい問題をしたい",
    "review": "授業の復習をしたい",
    "easy": "簡単な問題をしたい",
    "talk": "もっと先生と話したい",
    "homework": "学校の宿題をしたい",
}

MULTI_SEP = ", "


def join_choices(mapping: dict[str, str], keys: list[str]) -> str:
    return MULTI_SEP.join(mapping[key] for key in keys)


def scores(**overrides: int | None) -> dict[str, int | None]:
    values: dict[str, int | None] = {field: 5 for field in SCORE_FIELDS}
    values.update(overrides)
    return values


STUDENTS = [
    {
        "student_id": 101,
        "student_name": "青木はると",
        "grade": "中学２年",
        "case": "A",
        "case_label": "全体的に良好",
    },
    {
        "student_id": 102,
        "student_name": "加藤みお",
        "grade": "中学１年",
        "case": "B",
        "case_label": "授業満足度が低い",
    },
    {
        "student_id": 103,
        "student_name": "木村ゆうと",
        "grade": "中学３年",
        "case": "C",
        "case_label": "家庭学習に問題がある",
    },
    {
        "student_id": 104,
        "student_name": "斎藤あかり",
        "grade": "小学６年",
        "case": "D",
        "case_label": "学校関連が低い",
    },
    {
        "student_id": 105,
        "student_name": "高橋れん",
        "grade": "中学１年",
        "case": "E",
        "case_label": "勉強への気持ちが低下",
    },
    {
        "student_id": 106,
        "student_name": "中村さき",
        "grade": "中学２年",
        "case": "F",
        "case_label": "前月から急激に低下",
    },
    {
        "student_id": 107,
        "student_name": "林だいき",
        "grade": "小学５年",
        "case": "G",
        "case_label": "一見良好だが継続意向だけ低い",
    },
    {
        "student_id": 108,
        "student_name": "松本ひな",
        "grade": "小学４年",
        "case": "H",
        "case_label": "欠損データあり",
    },
    {
        "student_id": 109,
        "student_name": "山口かい",
        "grade": "中学３年",
        "case": "I",
        "case_label": "前月から改善",
    },
    {
        "student_id": 110,
        "student_name": "伊藤のあ",
        "grade": "小学６年",
        "case": "J",
        "case_label": "複数項目が悪化",
    },
]

MONTHS = [
    {
        "year_month": "2026-08",
        "responses": {
            101: {
                **scores(
                    juku_satisfaction=6,
                    juku_security=7,
                    juku_fun=6,
                    juku_clarity=7,
                    juku_question=7,
                    juku_motivation=6,
                    juku_homework=6,
                    juku_growth=6,
                    juku_continuation=7,
                    home_plan_achievement=6,
                    home_study_environment=7,
                    home_pace_drop=3,
                    home_problem_solving=6,
                    home_motivation=6,
                    home_growth=6,
                    home_continuation=6,
                    school_understanding=6,
                    school_security=7,
                    school_fun=6,
                    school_clarity=6,
                    school_question=6,
                    school_class_motivation=6,
                    school_class_satisfaction=6,
                    school_test_motivation=6,
                    school_test_satisfaction=6,
                    school_growth=6,
                    school_test_expectation=6,
                    school_continuation=6,
                    study_fun=6,
                    study_goal=7,
                    study_improvement_motivation=6,
                ),
                "timestamp": datetime(2026, 8, 12, 19, 20, 0),
                "study_daily_mitutes": STUDY_MINUTES["30_60"],
                "study_concerns": join_choices(CONCERNS, ["calc"]),
                "study_next_month_contents": join_choices(
                    NEXT_MONTH, ["math", "english", "hard"]
                ),
            },
            102: {
                **scores(
                    juku_satisfaction=3,
                    juku_security=5,
                    juku_fun=3,
                    juku_clarity=3,
                    juku_question=4,
                    juku_motivation=4,
                    juku_homework=4,
                    juku_growth=3,
                    juku_continuation=4,
                    home_plan_achievement=4,
                    home_study_environment=5,
                    home_pace_drop=5,
                    home_problem_solving=4,
                    home_motivation=4,
                    home_growth=4,
                    home_continuation=5,
                    school_understanding=5,
                    school_security=6,
                    school_fun=4,
                    school_clarity=4,
                    school_question=5,
                    school_class_motivation=5,
                    school_class_satisfaction=4,
                    school_test_motivation=5,
                    school_test_satisfaction=4,
                    school_growth=4,
                    school_test_expectation=4,
                    school_continuation=5,
                    study_fun=4,
                    study_goal=5,
                    study_improvement_motivation=4,
                ),
                "timestamp": datetime(2026, 8, 13, 18, 5, 0),
                "study_daily_mitutes": STUDY_MINUTES["10_30"],
                "study_concerns": join_choices(CONCERNS, ["motivation", "applied"]),
                "study_next_month_contents": join_choices(
                    NEXT_MONTH, ["review", "easy"]
                ),
            },
            103: {
                **scores(
                    juku_satisfaction=6,
                    juku_security=6,
                    juku_fun=6,
                    juku_clarity=6,
                    juku_question=6,
                    juku_motivation=5,
                    juku_homework=3,
                    juku_growth=5,
                    juku_continuation=6,
                    home_plan_achievement=2,
                    home_study_environment=4,
                    home_pace_drop=6,
                    home_problem_solving=3,
                    home_motivation=3,
                    home_growth=3,
                    home_continuation=2,
                    school_understanding=5,
                    school_security=6,
                    school_fun=5,
                    school_clarity=5,
                    school_question=4,
                    school_class_motivation=5,
                    school_class_satisfaction=4,
                    school_test_motivation=5,
                    school_test_satisfaction=4,
                    school_growth=4,
                    school_test_expectation=4,
                    school_continuation=5,
                    study_fun=5,
                    study_goal=6,
                    study_improvement_motivation=5,
                ),
                "timestamp": datetime(2026, 8, 14, 20, 10, 0),
                "study_daily_mitutes": STUDY_MINUTES["10_30"],
                "study_concerns": join_choices(CONCERNS, ["what", "how"]),
                "study_next_month_contents": join_choices(
                    NEXT_MONTH, ["homework", "review"]
                ),
            },
            104: {
                **scores(
                    juku_satisfaction=6,
                    juku_security=6,
                    juku_fun=6,
                    juku_clarity=6,
                    juku_question=6,
                    juku_motivation=6,
                    juku_homework=5,
                    juku_growth=5,
                    juku_continuation=6,
                    home_plan_achievement=5,
                    home_study_environment=6,
                    home_pace_drop=4,
                    home_problem_solving=5,
                    home_motivation=5,
                    home_growth=5,
                    home_continuation=5,
                    school_understanding=3,
                    school_security=4,
                    school_fun=3,
                    school_clarity=3,
                    school_question=3,
                    school_class_motivation=3,
                    school_class_satisfaction=3,
                    school_test_motivation=3,
                    school_test_satisfaction=2,
                    school_growth=3,
                    school_test_expectation=2,
                    school_continuation=4,
                    study_fun=4,
                    study_goal=5,
                    study_improvement_motivation=5,
                ),
                "timestamp": datetime(2026, 8, 15, 17, 40, 0),
                "study_daily_mitutes": STUDY_MINUTES["30_60"],
                "study_concerns": join_choices(CONCERNS, ["score", "applied"]),
                "study_next_month_contents": join_choices(
                    NEXT_MONTH, ["math", "science", "easy"]
                ),
            },
            105: {
                **scores(
                    juku_satisfaction=5,
                    juku_security=6,
                    juku_fun=5,
                    juku_clarity=5,
                    juku_question=5,
                    juku_motivation=4,
                    juku_homework=5,
                    juku_growth=4,
                    juku_continuation=5,
                    home_plan_achievement=4,
                    home_study_environment=5,
                    home_pace_drop=5,
                    home_problem_solving=4,
                    home_motivation=3,
                    home_growth=4,
                    home_continuation=4,
                    school_understanding=5,
                    school_security=6,
                    school_fun=4,
                    school_clarity=5,
                    school_question=5,
                    school_class_motivation=4,
                    school_class_satisfaction=4,
                    school_test_motivation=4,
                    school_test_satisfaction=4,
                    school_growth=4,
                    school_test_expectation=4,
                    school_continuation=5,
                    study_fun=2,
                    study_goal=2,
                    study_improvement_motivation=3,
                ),
                "timestamp": datetime(2026, 8, 16, 19, 5, 0),
                "study_daily_mitutes": STUDY_MINUTES["10_30"],
                "study_concerns": join_choices(CONCERNS, ["motivation", "what"]),
                "study_next_month_contents": join_choices(NEXT_MONTH, ["talk"]),
            },
            106: {
                **scores(
                    juku_satisfaction=6,
                    juku_security=6,
                    juku_fun=6,
                    juku_clarity=6,
                    juku_question=6,
                    juku_motivation=6,
                    juku_homework=6,
                    juku_growth=6,
                    juku_continuation=6,
                    home_plan_achievement=5,
                    home_study_environment=6,
                    home_pace_drop=4,
                    home_problem_solving=5,
                    home_motivation=5,
                    home_growth=5,
                    home_continuation=6,
                    school_understanding=6,
                    school_security=6,
                    school_fun=5,
                    school_clarity=6,
                    school_question=6,
                    school_class_motivation=6,
                    school_class_satisfaction=5,
                    school_test_motivation=6,
                    school_test_satisfaction=5,
                    school_growth=5,
                    school_test_expectation=6,
                    school_continuation=6,
                    study_fun=5,
                    study_goal=6,
                    study_improvement_motivation=6,
                ),
                "timestamp": datetime(2026, 8, 17, 20, 30, 0),
                "study_daily_mitutes": STUDY_MINUTES["30_60"],
                "study_concerns": join_choices(CONCERNS, ["calc"]),
                "study_next_month_contents": join_choices(
                    NEXT_MONTH, ["english", "math", "review"]
                ),
            },
            107: {
                **scores(
                    juku_satisfaction=7,
                    juku_security=7,
                    juku_fun=7,
                    juku_clarity=7,
                    juku_question=6,
                    juku_motivation=6,
                    juku_homework=6,
                    juku_growth=6,
                    juku_continuation=2,
                    home_plan_achievement=6,
                    home_study_environment=7,
                    home_pace_drop=3,
                    home_problem_solving=6,
                    home_motivation=6,
                    home_growth=6,
                    home_continuation=6,
                    school_understanding=6,
                    school_security=7,
                    school_fun=6,
                    school_clarity=6,
                    school_question=6,
                    school_class_motivation=6,
                    school_class_satisfaction=6,
                    school_test_motivation=6,
                    school_test_satisfaction=6,
                    school_growth=6,
                    school_test_expectation=6,
                    school_continuation=6,
                    study_fun=6,
                    study_goal=7,
                    study_improvement_motivation=6,
                ),
                "timestamp": datetime(2026, 8, 18, 18, 15, 0),
                "study_daily_mitutes": STUDY_MINUTES["30_60"],
                "study_concerns": join_choices(CONCERNS, ["calc"]),
                "study_next_month_contents": join_choices(
                    NEXT_MONTH, ["kanji", "math", "hard"]
                ),
            },
            108: {
                **scores(
                    juku_satisfaction=5,
                    juku_security=6,
                    juku_fun=5,
                    juku_clarity=None,
                    juku_question=5,
                    juku_motivation=5,
                    juku_homework=5,
                    juku_growth=None,
                    juku_continuation=5,
                    home_plan_achievement=4,
                    home_study_environment=5,
                    home_pace_drop=5,
                    home_problem_solving=None,
                    home_motivation=4,
                    home_growth=None,
                    home_continuation=5,
                    school_understanding=5,
                    school_security=6,
                    school_fun=5,
                    school_clarity=5,
                    school_question=None,
                    school_class_motivation=5,
                    school_class_satisfaction=4,
                    school_test_motivation=5,
                    school_test_satisfaction=4,
                    school_growth=5,
                    school_test_expectation=5,
                    school_continuation=5,
                    study_fun=5,
                    study_goal=None,
                    study_improvement_motivation=5,
                ),
                "timestamp": datetime(2026, 8, 19, 19, 50, 0),
                "study_daily_mitutes": STUDY_MINUTES["10_30"],
                "study_concerns": join_choices(CONCERNS, ["score"]),
                "study_next_month_contents": join_choices(
                    NEXT_MONTH, ["reading", "kanji"]
                ),
            },
            109: {
                **scores(
                    juku_satisfaction=3,
                    juku_security=5,
                    juku_fun=3,
                    juku_clarity=4,
                    juku_question=4,
                    juku_motivation=3,
                    juku_homework=3,
                    juku_growth=3,
                    juku_continuation=4,
                    home_plan_achievement=3,
                    home_study_environment=5,
                    home_pace_drop=5,
                    home_problem_solving=3,
                    home_motivation=3,
                    home_growth=3,
                    home_continuation=3,
                    school_understanding=4,
                    school_security=5,
                    school_fun=4,
                    school_clarity=4,
                    school_question=4,
                    school_class_motivation=4,
                    school_class_satisfaction=3,
                    school_test_motivation=3,
                    school_test_satisfaction=3,
                    school_growth=3,
                    school_test_expectation=3,
                    school_continuation=4,
                    study_fun=3,
                    study_goal=4,
                    study_improvement_motivation=4,
                ),
                "timestamp": datetime(2026, 8, 20, 20, 5, 0),
                "study_daily_mitutes": STUDY_MINUTES["10_30"],
                "study_concerns": join_choices(CONCERNS, ["motivation", "how", "score"]),
                "study_next_month_contents": join_choices(
                    NEXT_MONTH, ["review", "easy", "talk"]
                ),
            },
            110: {
                **scores(
                    juku_satisfaction=5,
                    juku_security=6,
                    juku_fun=5,
                    juku_clarity=5,
                    juku_question=5,
                    juku_motivation=5,
                    juku_homework=5,
                    juku_growth=5,
                    juku_continuation=5,
                    home_plan_achievement=5,
                    home_study_environment=5,
                    home_pace_drop=4,
                    home_problem_solving=5,
                    home_motivation=5,
                    home_growth=5,
                    home_continuation=5,
                    school_understanding=5,
                    school_security=6,
                    school_fun=5,
                    school_clarity=5,
                    school_question=5,
                    school_class_motivation=5,
                    school_class_satisfaction=5,
                    school_test_motivation=5,
                    school_test_satisfaction=5,
                    school_growth=5,
                    school_test_expectation=5,
                    school_continuation=5,
                    study_fun=5,
                    study_goal=5,
                    study_improvement_motivation=5,
                ),
                "timestamp": datetime(2026, 8, 21, 18, 45, 0),
                "study_daily_mitutes": STUDY_MINUTES["30_60"],
                "study_concerns": join_choices(CONCERNS, ["calc", "applied"]),
                "study_next_month_contents": join_choices(
                    NEXT_MONTH, ["math", "english", "social"]
                ),
            },
        },
    },
    {
        "year_month": "2026-09",
        "responses": {
            101: {
                **scores(
                    juku_satisfaction=7,
                    juku_security=7,
                    juku_fun=6,
                    juku_clarity=7,
                    juku_question=7,
                    juku_motivation=7,
                    juku_homework=6,
                    juku_growth=6,
                    juku_continuation=7,
                    home_plan_achievement=6,
                    home_study_environment=7,
                    home_pace_drop=3,
                    home_problem_solving=6,
                    home_motivation=6,
                    home_growth=6,
                    home_continuation=7,
                    school_understanding=6,
                    school_security=7,
                    school_fun=6,
                    school_clarity=6,
                    school_question=6,
                    school_class_motivation=6,
                    school_class_satisfaction=6,
                    school_test_motivation=7,
                    school_test_satisfaction=6,
                    school_growth=6,
                    school_test_expectation=6,
                    school_continuation=6,
                    study_fun=6,
                    study_goal=7,
                    study_improvement_motivation=7,
                ),
                "timestamp": datetime(2026, 9, 10, 19, 10, 0),
                "study_daily_mitutes": STUDY_MINUTES["30_60"],
                "study_concerns": join_choices(CONCERNS, ["calc"]),
                "study_next_month_contents": join_choices(
                    NEXT_MONTH, ["math", "english", "hard"]
                ),
            },
            102: {
                **scores(
                    juku_satisfaction=3,
                    juku_security=5,
                    juku_fun=3,
                    juku_clarity=2,
                    juku_question=3,
                    juku_motivation=3,
                    juku_homework=4,
                    juku_growth=3,
                    juku_continuation=3,
                    home_plan_achievement=4,
                    home_study_environment=5,
                    home_pace_drop=5,
                    home_problem_solving=4,
                    home_motivation=3,
                    home_growth=3,
                    home_continuation=4,
                    school_understanding=5,
                    school_security=6,
                    school_fun=4,
                    school_clarity=4,
                    school_question=4,
                    school_class_motivation=4,
                    school_class_satisfaction=4,
                    school_test_motivation=4,
                    school_test_satisfaction=4,
                    school_growth=4,
                    school_test_expectation=3,
                    school_continuation=4,
                    study_fun=3,
                    study_goal=5,
                    study_improvement_motivation=4,
                ),
                "timestamp": datetime(2026, 9, 11, 18, 25, 0),
                "study_daily_mitutes": STUDY_MINUTES["10_30"],
                "study_concerns": join_choices(
                    CONCERNS, ["motivation", "applied", "score"]
                ),
                "study_next_month_contents": join_choices(
                    NEXT_MONTH, ["review", "easy", "talk"]
                ),
            },
            103: {
                **scores(
                    juku_satisfaction=6,
                    juku_security=6,
                    juku_fun=6,
                    juku_clarity=6,
                    juku_question=6,
                    juku_motivation=5,
                    juku_homework=3,
                    juku_growth=5,
                    juku_continuation=6,
                    home_plan_achievement=2,
                    home_study_environment=4,
                    home_pace_drop=7,
                    home_problem_solving=3,
                    home_motivation=2,
                    home_growth=3,
                    home_continuation=2,
                    school_understanding=4,
                    school_security=6,
                    school_fun=5,
                    school_clarity=4,
                    school_question=4,
                    school_class_motivation=5,
                    school_class_satisfaction=4,
                    school_test_motivation=5,
                    school_test_satisfaction=3,
                    school_growth=4,
                    school_test_expectation=4,
                    school_continuation=5,
                    study_fun=4,
                    study_goal=6,
                    study_improvement_motivation=5,
                ),
                "timestamp": datetime(2026, 9, 12, 20, 40, 0),
                "study_daily_mitutes": STUDY_MINUTES["10_30"],
                "study_concerns": join_choices(CONCERNS, ["what", "how", "calc"]),
                "study_next_month_contents": join_choices(
                    NEXT_MONTH, ["homework", "review", "easy"]
                ),
            },
            104: {
                **scores(
                    juku_satisfaction=6,
                    juku_security=6,
                    juku_fun=6,
                    juku_clarity=6,
                    juku_question=6,
                    juku_motivation=6,
                    juku_homework=5,
                    juku_growth=5,
                    juku_continuation=6,
                    home_plan_achievement=5,
                    home_study_environment=6,
                    home_pace_drop=4,
                    home_problem_solving=5,
                    home_motivation=5,
                    home_growth=5,
                    home_continuation=5,
                    school_understanding=2,
                    school_security=3,
                    school_fun=3,
                    school_clarity=2,
                    school_question=3,
                    school_class_motivation=3,
                    school_class_satisfaction=2,
                    school_test_motivation=3,
                    school_test_satisfaction=2,
                    school_growth=2,
                    school_test_expectation=2,
                    school_continuation=3,
                    study_fun=4,
                    study_goal=5,
                    study_improvement_motivation=5,
                ),
                "timestamp": datetime(2026, 9, 13, 17, 15, 0),
                "study_daily_mitutes": STUDY_MINUTES["30_60"],
                "study_concerns": join_choices(CONCERNS, ["score", "applied", "what"]),
                "study_next_month_contents": join_choices(
                    NEXT_MONTH, ["math", "science", "easy", "review"]
                ),
            },
            105: {
                **scores(
                    juku_satisfaction=5,
                    juku_security=6,
                    juku_fun=5,
                    juku_clarity=5,
                    juku_question=5,
                    juku_motivation=4,
                    juku_homework=4,
                    juku_growth=4,
                    juku_continuation=4,
                    home_plan_achievement=4,
                    home_study_environment=5,
                    home_pace_drop=5,
                    home_problem_solving=4,
                    home_motivation=3,
                    home_growth=3,
                    home_continuation=3,
                    school_understanding=5,
                    school_security=6,
                    school_fun=4,
                    school_clarity=5,
                    school_question=4,
                    school_class_motivation=4,
                    school_class_satisfaction=4,
                    school_test_motivation=3,
                    school_test_satisfaction=3,
                    school_growth=4,
                    school_test_expectation=3,
                    school_continuation=4,
                    study_fun=2,
                    study_goal=2,
                    study_improvement_motivation=2,
                ),
                "timestamp": datetime(2026, 9, 14, 19, 35, 0),
                "study_daily_mitutes": STUDY_MINUTES["10_30"],
                "study_concerns": join_choices(CONCERNS, ["motivation", "what", "how"]),
                "study_next_month_contents": join_choices(NEXT_MONTH, ["talk", "easy"]),
            },
            106: {
                **scores(
                    juku_satisfaction=3,
                    juku_security=5,
                    juku_fun=3,
                    juku_clarity=4,
                    juku_question=4,
                    juku_motivation=3,
                    juku_homework=4,
                    juku_growth=3,
                    juku_continuation=4,
                    home_plan_achievement=4,
                    home_study_environment=6,
                    home_pace_drop=5,
                    home_problem_solving=4,
                    home_motivation=3,
                    home_growth=3,
                    home_continuation=4,
                    school_understanding=5,
                    school_security=6,
                    school_fun=4,
                    school_clarity=5,
                    school_question=5,
                    school_class_motivation=4,
                    school_class_satisfaction=4,
                    school_test_motivation=4,
                    school_test_satisfaction=4,
                    school_growth=4,
                    school_test_expectation=3,
                    school_continuation=5,
                    study_fun=3,
                    study_goal=5,
                    study_improvement_motivation=4,
                ),
                "timestamp": datetime(2026, 9, 15, 20, 20, 0),
                "study_daily_mitutes": STUDY_MINUTES["10_30"],
                "study_concerns": join_choices(CONCERNS, ["motivation", "score"]),
                "study_next_month_contents": join_choices(
                    NEXT_MONTH, ["review", "talk"]
                ),
            },
            107: {
                **scores(
                    juku_satisfaction=7,
                    juku_security=7,
                    juku_fun=6,
                    juku_clarity=7,
                    juku_question=6,
                    juku_motivation=6,
                    juku_homework=6,
                    juku_growth=6,
                    juku_continuation=2,
                    home_plan_achievement=6,
                    home_study_environment=7,
                    home_pace_drop=3,
                    home_problem_solving=6,
                    home_motivation=6,
                    home_growth=6,
                    home_continuation=6,
                    school_understanding=6,
                    school_security=7,
                    school_fun=6,
                    school_clarity=6,
                    school_question=6,
                    school_class_motivation=6,
                    school_class_satisfaction=6,
                    school_test_motivation=6,
                    school_test_satisfaction=6,
                    school_growth=6,
                    school_test_expectation=6,
                    school_continuation=6,
                    study_fun=6,
                    study_goal=7,
                    study_improvement_motivation=6,
                ),
                "timestamp": datetime(2026, 9, 16, 18, 55, 0),
                "study_daily_mitutes": STUDY_MINUTES["30_60"],
                "study_concerns": join_choices(CONCERNS, ["calc"]),
                "study_next_month_contents": join_choices(
                    NEXT_MONTH, ["kanji", "math", "hard"]
                ),
            },
            108: {
                **scores(
                    juku_satisfaction=5,
                    juku_security=6,
                    juku_fun=None,
                    juku_clarity=5,
                    juku_question=5,
                    juku_motivation=5,
                    juku_homework=None,
                    juku_growth=5,
                    juku_continuation=5,
                    home_plan_achievement=4,
                    home_study_environment=None,
                    home_pace_drop=5,
                    home_problem_solving=4,
                    home_motivation=4,
                    home_growth=4,
                    home_continuation=5,
                    school_understanding=5,
                    school_security=6,
                    school_fun=5,
                    school_clarity=None,
                    school_question=5,
                    school_class_motivation=5,
                    school_class_satisfaction=4,
                    school_test_motivation=5,
                    school_test_satisfaction=None,
                    school_growth=5,
                    school_test_expectation=5,
                    school_continuation=5,
                    study_fun=5,
                    study_goal=5,
                    study_improvement_motivation=None,
                ),
                "timestamp": datetime(2026, 9, 17, 19, 5, 0),
                "study_daily_mitutes": STUDY_MINUTES["10_30"],
                "study_concerns": join_choices(CONCERNS, ["score", "calc"]),
                "study_next_month_contents": join_choices(
                    NEXT_MONTH, ["reading", "kanji", "homework"]
                ),
            },
            109: {
                **scores(
                    juku_satisfaction=6,
                    juku_security=6,
                    juku_fun=5,
                    juku_clarity=6,
                    juku_question=6,
                    juku_motivation=5,
                    juku_homework=5,
                    juku_growth=5,
                    juku_continuation=6,
                    home_plan_achievement=5,
                    home_study_environment=6,
                    home_pace_drop=4,
                    home_problem_solving=5,
                    home_motivation=5,
                    home_growth=5,
                    home_continuation=5,
                    school_understanding=5,
                    school_security=6,
                    school_fun=5,
                    school_clarity=5,
                    school_question=5,
                    school_class_motivation=5,
                    school_class_satisfaction=5,
                    school_test_motivation=5,
                    school_test_satisfaction=5,
                    school_growth=5,
                    school_test_expectation=5,
                    school_continuation=5,
                    study_fun=5,
                    study_goal=6,
                    study_improvement_motivation=6,
                ),
                "timestamp": datetime(2026, 9, 18, 20, 15, 0),
                "study_daily_mitutes": STUDY_MINUTES["30_60"],
                "study_concerns": join_choices(CONCERNS, ["how"]),
                "study_next_month_contents": join_choices(
                    NEXT_MONTH, ["review", "english", "math"]
                ),
            },
            110: {
                **scores(
                    juku_satisfaction=4,
                    juku_security=5,
                    juku_fun=4,
                    juku_clarity=4,
                    juku_question=4,
                    juku_motivation=4,
                    juku_homework=4,
                    juku_growth=3,
                    juku_continuation=4,
                    home_plan_achievement=4,
                    home_study_environment=5,
                    home_pace_drop=5,
                    home_problem_solving=4,
                    home_motivation=3,
                    home_growth=3,
                    home_continuation=4,
                    school_understanding=4,
                    school_security=5,
                    school_fun=4,
                    school_clarity=4,
                    school_question=4,
                    school_class_motivation=4,
                    school_class_satisfaction=3,
                    school_test_motivation=4,
                    school_test_satisfaction=3,
                    school_growth=3,
                    school_test_expectation=3,
                    school_continuation=4,
                    study_fun=3,
                    study_goal=4,
                    study_improvement_motivation=4,
                ),
                "timestamp": datetime(2026, 9, 19, 18, 30, 0),
                "study_daily_mitutes": STUDY_MINUTES["10_30"],
                "study_concerns": join_choices(
                    CONCERNS, ["calc", "applied", "motivation"]
                ),
                "study_next_month_contents": join_choices(
                    NEXT_MONTH, ["math", "english", "easy", "review"]
                ),
            },
        },
    },
]


def build_rows() -> list[dict]:
    students_by_id = {student["student_id"]: student for student in STUDENTS}
    rows: list[dict] = []
    for month in MONTHS:
        for student_id, response in month["responses"].items():
            student = students_by_id[student_id]
            row = {header: None for header in HEADERS}
            row.update(response)
            row["student_id"] = student["student_id"]
            row["student_name"] = student["student_name"]
            row["grade"] = student["grade"]
            row["_case"] = student["case"]
            row["_case_label"] = student["case_label"]
            row["_year_month"] = month["year_month"]
            rows.append(row)
    rows.sort(key=lambda item: (item["timestamp"], item["student_id"]))
    return rows


def cell_value(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    return value


def write_csv(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=HEADERS, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({header: cell_value(row[header]) for header in HEADERS})


def write_json(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "source": "demo",
        "note": "共有テストデータと同じ列構造のデモ用データ。実在の生徒情報は含まない。",
        "months": ["2026-08", "2026-09"],
        "students": [
            {
                "student_id": student["student_id"],
                "student_name": student["student_name"],
                "grade": student["grade"],
                "case": student["case"],
                "case_label": student["case_label"],
            }
            for student in STUDENTS
        ],
        "responses": [
            {
                **{header: cell_value(row[header]) if row[header] is not None else None for header in HEADERS},
                "case": row["_case"],
                "case_label": row["_case_label"],
                "year_month": row["_year_month"],
            }
            for row in rows
        ],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_xlsx(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "demo_survey_responses"

    header_font = Font(bold=True)
    for col, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=1, column=col, value=header)
        cell.font = header_font

    for row_index, row in enumerate(rows, start=2):
        for col, header in enumerate(HEADERS, start=1):
            value = row[header]
            cell = sheet.cell(row=row_index, column=col, value=value)
            if header == "timestamp" and isinstance(value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            if header in SCORE_FIELDS and value is not None:
                cell.alignment = Alignment(horizontal="center")

    last_column = get_column_letter(len(HEADERS))
    last_row = len(rows) + 1
    table = Table(displayName="demo_survey_responses", ref=f"A1:{last_column}{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)

    widths = {
        "timestamp": 22,
        "student_name": 14,
        "student_id": 12,
        "grade": 12,
        "study_daily_mitutes": 18,
        "study_concerns": 56,
        "study_next_month_contents": 64,
    }
    for col, header in enumerate(HEADERS, start=1):
        sheet.column_dimensions[get_column_letter(col)].width = widths.get(header, 16)

    workbook.save(path)


def main() -> None:
    project_root = Path(__file__).resolve().parents[1]
    output_dir = project_root / "data" / "demo"
    rows = build_rows()

    csv_path = output_dir / "demo_survey_responses.csv"
    json_path = output_dir / "demo_survey_responses.json"
    xlsx_path = output_dir / "demo_survey_responses.xlsx"

    write_csv(csv_path, rows)
    write_json(json_path, rows)
    write_xlsx(xlsx_path, rows)

    print(f"rows={len(rows)}")
    print(f"students={len(STUDENTS)}")
    print(f"months={', '.join(month['year_month'] for month in MONTHS)}")
    print(f"wrote {csv_path}")
    print(f"wrote {json_path}")
    print(f"wrote {xlsx_path}")


if __name__ == "__main__":
    main()
